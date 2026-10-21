# -*- coding: utf-8 -*-
"""Importar el personal desde un Excel (la "BASE PEYMAR" de RH) a RH sin Microsip.

- Lee la hoja (por default la que se llame "BASE PEYMAR" o la primera que tenga
  la columna "NOMBRE COMPLETO") y reconoce las columnas por su encabezado.
- Cada renglón se compara con el personal que ya existe (por CURP o por nombre)
  y con la gente que ya está en los checadores (por nombre: el checador guarda
  hasta 24 letras y a veces en otro orden), para usar el MISMO No. de empleado
  que ya tiene en el biométrico y no duplicar a nadie.
- Primero se hace la vista previa (plan); al confirmar se aplica.
Aquí solo está la lógica pura (sin base de datos) para poder probarla.
"""
import io
import re
import unicodedata
from collections import Counter
from datetime import date, datetime

# encabezado normalizado → campo
COLUMNAS = {
    "NOMBRE COMPLETO": "nombre_completo", "NOMBRE": "nombre_completo",
    "FECHA DE INGRESO": "fecha_ingreso", "PUESTO": "puesto", "AREA": "departamento", "DEPARTAMENTO": "departamento",
    "JEFE DIRECTO": "jefe_directo", "TELEFONO": "telefono_whatsapp", "TELEFONOS": "telefono_whatsapp",
    "CORREO ELECTRONICO": "email", "E-MAIL": "email", "EMAIL": "email", "CORREO": "email",
    "NSS": "nss", "RFC": "rfc", "CURP": "curp", "NUMERO DE INE": "numero_ine", "INE": "numero_ine",
    "FECHA DE NACIMIENTO": "fecha_nacimiento", "F. NACIMIENTO": "fecha_nacimiento",
    "EST.CIVIL": "estado_civil", "ESTADO CIVIL": "estado_civil",
    "SUELDO SEMANAL": "sueldo_semanal", "BONO": "bono_semanal", "SEGURO IMSS": "seguro_imss",
    "FECHA DE ALTA": "fecha_alta_imss", "RAZON SOCIAL": "razon_social",
    "NO. DE EMPLEADO": "numero_empleado", "NO. EMPLEADO": "numero_empleado", "NUMERO DE EMPLEADO": "numero_empleado",
    "NO EMPLEADO": "numero_empleado",
}
VACIOS = {"", "-", "--", "N/A", "NA", "NONE", "S/N", "0"}


def sin_acentos(t):
    return "".join(c for c in unicodedata.normalize("NFD", str(t)) if unicodedata.category(c) != "Mn")


def norm(t):
    """'José  Pérez ' → 'JOSE PEREZ' (para comparar)."""
    t = sin_acentos(t or "").upper()
    t = re.sub(r"[^A-Z0-9Ñ ]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def tokens(t):
    return [x for x in norm(t).split(" ") if x]


def _texto(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = re.sub(r"\s+", " ", str(v)).strip()
    return None if t.upper() in VACIOS else t


def _fecha(v):
    if isinstance(v, datetime):
        d = v.date()
    elif isinstance(v, date):
        d = v
    else:
        t = _texto(v)
        if not t:
            return None
        m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})$", t)
        m2 = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", t)
        try:
            if m:
                a = int(m.group(3)); a = a + 2000 if a < 100 else a
                d = date(a, int(m.group(2)), int(m.group(1)))
            elif m2:
                d = date(int(m2.group(1)), int(m2.group(2)), int(m2.group(3)))
            else:
                return None
        except ValueError:
            return None
    return d.isoformat() if 1930 <= d.year <= 2100 else None


def _numero(v):
    if v is None or (isinstance(v, str) and _texto(v) is None):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = re.sub(r"[^0-9.\-]", "", str(v))
    try:
        return float(t) if t not in ("", "-", ".") else None
    except ValueError:
        return None


def _digitos(v):
    if v is None:
        return None
    t = re.sub(r"\D", "", str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
    return t or None


def limpiar(campo, v):
    if campo in ("fecha_ingreso", "fecha_nacimiento", "fecha_alta_imss"):
        return _fecha(v)
    if campo in ("sueldo_semanal", "bono_semanal"):
        n = _numero(v)
        return round(n, 2) if n is not None and n >= 0 else None
    if campo == "nss":
        d = _digitos(v)
        if d and len(d) == 10:
            d = "0" + d  # el Excel le quita el 0 de enfrente
        return d
    if campo == "telefono_whatsapp":
        d = _digitos(v)
        return d[-10:] if d and len(d) >= 10 else d
    if campo == "numero_empleado":
        d = _digitos(v)
        return d.lstrip("0") or None if d else None
    t = _texto(v)
    if t is None:
        return None
    if campo in ("rfc", "curp", "numero_ine"):
        return re.sub(r"\s+", "", t).upper()
    if campo == "email":
        t = t.replace(" ", "").lower()
        return t if "@" in t else None
    if campo == "estado_civil":
        return norm(t).capitalize()
    if campo in ("nombre_completo", "puesto", "departamento", "jefe_directo", "razon_social", "seguro_imss"):
        return t.upper() if campo != "puesto" else t.upper()
    return t


def leer_excel(contenido: bytes, hoja=None):
    """→ (filas [{campo: valor, '_fila': n}], hoja_usada, hojas_disponibles, columnas_reconocidas)."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    candidatas = []
    for ws in wb.worksheets:
        enc = None
        for i, r in enumerate(ws.iter_rows(values_only=True, max_row=10, max_col=80)):
            nombres = [norm(x) for x in r]
            if "NOMBRE COMPLETO" in nombres:
                enc = (i, nombres)
                break
        if enc:
            candidatas.append((ws, enc))
    if not candidatas:
        raise ValueError("No encontré una hoja con la columna «NOMBRE COMPLETO».")
    hojas = [ws.title for ws, _ in candidatas]
    elegida = None
    if hoja:
        elegida = next((c for c in candidatas if c[0].title == hoja), None)
    if not elegida:
        elegida = next((c for c in candidatas if "PEYMAR" in norm(c[0].title)), candidatas[0])
    ws, (i_enc, nombres) = elegida
    mapa = {}
    for j, h in enumerate(nombres):
        campo = COLUMNAS.get(h)
        if campo and campo not in mapa.values():
            mapa[j] = campo
    filas = []
    j_nombre = next(j for j, c in mapa.items() if c == "nombre_completo")
    vacios = 0
    for n, r in enumerate(ws.iter_rows(values_only=True, min_row=i_enc + 2, max_col=max(mapa) + 1 if mapa else 1), start=i_enc + 2):
        # Muchas bases traen miles de renglones vacíos con formato: después de 25 seguidos, se acaba.
        if j_nombre >= len(r) or _texto(r[j_nombre]) is None:
            vacios += 1
            if vacios >= 25:
                break
            continue
        vacios = 0
        datos = {"_fila": n}
        for j, campo in mapa.items():
            if j < len(r):
                v = limpiar(campo, r[j])
                if v not in (None, ""):
                    datos[campo] = v
        if datos.get("nombre_completo") and len(tokens(datos["nombre_completo"])) >= 2:
            filas.append(datos)
    return filas, ws.title, hojas, sorted(set(mapa.values()))


# ---------------------------------------------------------------- coincidencias

def _coincide_checador(nombre_chk, toks_persona):
    """El nombre del checador (hasta 24 letras, en cualquier orden, la última palabra
    puede venir cortada) contra las palabras del nombre completo."""
    tk = tokens(nombre_chk)
    if len(tk) < 2:
        return False
    disponibles = list(toks_persona)
    for k, t in enumerate(tk):
        ultima = k == len(tk) - 1
        hallado = next((p for p in disponibles if p == t or (ultima and len(t) >= 2 and p.startswith(t))), None)
        if not hallado:
            return False
        disponibles.remove(hallado)
    return True


def _misma_persona(nombre_a, nombre_b):
    a, b = sorted(tokens(nombre_a)), sorted(tokens(nombre_b))
    return a == b and len(a) >= 2


def planear(filas, personal, gente, siguiente_numero, departamentos):
    """filas: del Excel. personal: users de la empresa (activos y bajas).
    gente: rh_checador_gente activa [{sn, checador, codigo, nombre_checador, nombre, usuario_id}].
    → {'renglones': [...], 'departamentos_nuevos': [...], 'checador_sin_persona': [...], 'totales': {...}}"""
    deptos_norm = {norm(d): d for d in departamentos}
    deptos_nuevos = {}
    usados = {str(p.get("numero_empleado")).lstrip("0") for p in personal
              if p.get("activo") and str(p.get("numero_empleado") or "").strip()}
    por_curp = {p["curp"].upper(): p for p in personal if p.get("curp")}
    renglones, ya_persona = [], set()

    # 1) A quién del checador corresponde cada renglón.
    candidatos_codigo = {}  # (sn, codigo) -> [indices de renglón]
    for i, f in enumerate(filas):
        tk = tokens(f["nombre_completo"])
        for g in gente:
            nom = g.get("nombre_checador") or g.get("nombre")
            if nom and _coincide_checador(nom, tk):
                candidatos_codigo.setdefault((g["sn"], g["codigo"]), []).append(i)
    asignados = {}  # i -> [g]
    ambiguos = set()
    for (sn, cod), idx in candidatos_codigo.items():
        g = next(x for x in gente if x["sn"] == sn and x["codigo"] == cod)
        if len(idx) == 1:
            asignados.setdefault(idx[0], []).append(g)
        else:
            ambiguos.add((sn, cod))

    numero_sig = int(siguiente_numero or 1)
    for i, f in enumerate(filas):
        avisos = []
        existente = None
        if f.get("curp") and f["curp"] in por_curp:
            existente = por_curp[f["curp"]]
        if not existente:
            iguales = [p for p in personal if _misma_persona(p["nombre_completo"], f["nombre_completo"])]
            if len(iguales) == 1:
                existente = iguales[0]
            elif len(iguales) > 1:
                activos = [p for p in iguales if p.get("activo")]
                existente = activos[0] if len(activos) == 1 else None
                if not existente:
                    avisos.append("Hay más de una persona con ese nombre en el sistema: se da de alta aparte, revísalo.")
        if existente and existente["id"] in ya_persona:
            avisos.append("Está repetido en el Excel: solo se toma el primero.")
            renglones.append({"fila": f["_fila"], "nombre": f["nombre_completo"], "accion": "omitir", "avisos": avisos})
            continue
        if existente:
            ya_persona.add(existente["id"])

        # Checadores: lo que ya está ligado a esa persona manda; si no, por nombre.
        en_chk = [g for g in gente if existente and g.get("usuario_id") == existente["id"]]
        por_nombre = [g for g in asignados.get(i, [])
                      if not g.get("usuario_id") or (existente and g["usuario_id"] == existente["id"])]
        ligados_a_otro = [g for g in asignados.get(i, []) if g.get("usuario_id") and (not existente or g["usuario_id"] != existente["id"])]
        if ligados_a_otro:
            avisos.append("En el checador hay alguien con un nombre parecido pero ya ligado a otra persona: " +
                          ", ".join(f"{g['checador']} No. {g['codigo']}" for g in ligados_a_otro))
        ligar = [g for g in por_nombre if not g.get("usuario_id")]
        todos_chk = en_chk + [g for g in ligar if g not in en_chk]
        codigos = Counter(str(g["codigo"]).lstrip("0") for g in todos_chk)

        numero, origen = None, None
        actual = str((existente or {}).get("numero_empleado") or "").strip().lstrip("0")
        if actual:
            numero, origen = actual, "ya_tenia"
        elif f.get("numero_empleado") and f["numero_empleado"] not in usados:
            numero, origen = f["numero_empleado"], "excel"
        elif codigos:
            cod = codigos.most_common(1)[0][0]
            if cod in usados:
                avisos.append(f"En el checador tiene el No. {cod}, pero ese número ya lo tiene otra persona del sistema.")
            else:
                numero, origen = cod, "checador"
        if not numero:
            while str(numero_sig) in usados or any(str(g["codigo"]).lstrip("0") == str(numero_sig) for g in gente):
                numero_sig += 1
            numero, origen = str(numero_sig), "consecutivo"
            numero_sig += 1
        usados.add(numero)
        distintos = [g for g in todos_chk if str(g["codigo"]).lstrip("0") != numero]
        if distintos:
            avisos.append("En " + ", ".join(f"{g['checador']} tiene el No. {g['codigo']}" for g in distintos) +
                          f": para dejarle el {numero} usa «Poner el No. de empleado» en Conexión checadores.")
        ligar = [g for g in ligar if str(g["codigo"]).lstrip("0") == numero]

        depto = f.get("departamento")
        if depto:
            k = norm(depto)
            if k in deptos_norm:
                depto = deptos_norm[k]
            else:
                deptos_nuevos.setdefault(k, depto)
                depto = deptos_nuevos[k]
        if not f.get("fecha_ingreso"):
            avisos.append("Sin fecha de ingreso: no se le calculan vacaciones hasta que la pongas.")
        elif f["fecha_ingreso"] > date.today().isoformat():
            avisos.append("La fecha de ingreso es futura: no se puso.")
            f = {**f}; f.pop("fecha_ingreso")
        datos = {k: v for k, v in f.items() if not k.startswith("_") and k != "numero_empleado"}
        if depto:
            datos["departamento"] = depto
        cambios = []
        if existente:
            datos.pop("nombre_completo", None)  # ya existe: no se le cambia el nombre
            for k, v in datos.items():
                a = existente.get(k)
                if a in (None, "") or str(a).strip().upper() != str(v).strip().upper():
                    if not (isinstance(a, (int, float)) and isinstance(v, (int, float)) and abs(float(a) - float(v)) < 0.005):
                        cambios.append(k)
            if not existente.get("activo"):
                avisos.append("En el sistema está dado de BAJA: se actualizan sus datos pero no se reactiva.")
        renglones.append({
            "fila": f["_fila"], "nombre": f["nombre_completo"],
            "accion": "actualizar" if existente else "nuevo",
            "usuario_id": existente["id"] if existente else None,
            "numero": numero, "origen_numero": origen,
            "checadores": sorted({g["checador"] for g in todos_chk if str(g["codigo"]).lstrip("0") == numero}),
            "ligar": [{"sn": g["sn"], "codigo": g["codigo"]} for g in ligar],
            "datos": datos, "cambios": cambios if existente else sorted(datos), "avisos": avisos,
        })
    usados_chk = {(l["sn"], l["codigo"]) for r in renglones for l in r.get("ligar", [])}
    sin_persona = [{"checador": g["checador"], "codigo": g["codigo"], "nombre": g.get("nombre_checador") or g.get("nombre"),
                    "ambiguo": (g["sn"], g["codigo"]) in ambiguos}
                   for g in gente if not g.get("usuario_id") and (g["sn"], g["codigo"]) not in usados_chk]
    tot = Counter(r["accion"] for r in renglones)
    return {
        "renglones": renglones,
        "departamentos_nuevos": sorted(deptos_nuevos.values()),
        "checador_sin_persona": sorted(sin_persona, key=lambda x: (x["checador"], len(x["codigo"]), x["codigo"])),
        "totales": {
            "renglones": len(renglones), "nuevos": tot.get("nuevo", 0), "actualizar": tot.get("actualizar", 0),
            "omitidos": tot.get("omitir", 0),
            "con_checador": sum(1 for r in renglones if r.get("checadores")),
            "sin_checador": sum(1 for r in renglones if r["accion"] != "omitir" and not r.get("checadores")),
            "con_avisos": sum(1 for r in renglones if r["avisos"]),
        },
    }
