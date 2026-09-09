"""Integración con la API de Claude (Anthropic) para leer fotos/imágenes
Y AHORA TAMBIÉN documentos (PDF, Word, Excel, CSV) con listas de
artículos, extraer nombre + cantidad de cada uno, para después
buscarlos en Microsip y armar una cotización.

Requiere la variable de entorno ANTHROPIC_API_KEY. Si no está
configurada, las funciones de este módulo lanzan RuntimeError con un
mensaje claro (la app sigue funcionando normal para todo lo demás).
"""
import base64
import io
import json
import os
import re

import requests

import db

ANTHROPIC_API_URL = "https://api.anthropic.com/v1/messages"
ANTHROPIC_VERSION = "2023-06-01"
# Sonnet da mejor lectura de letra manuscrita/fotos de empaque y de
# documentos largos que un modelo más chico — el costo por archivo
# sigue siendo unos centavos.
MODELO_LECTURA_IMAGEN = "claude-sonnet-5"

PROMPT_LECTURA_LISTA = """Este archivo puede ser: una lista escrita a mano, una lista impresa o \
captura de pantalla, fotos de productos físicos (empaques/etiquetas), o un documento \
(PDF/Word/Excel) con una lista o tabla de artículos/pedido/cotización. \
Identifica cada artículo/producto que aparezca y su cantidad.

Reglas:
- Si no hay cantidad especificada para un artículo, usa 1.
- Si es una foto de un producto físico, usa el nombre/marca/modelo que \
alcances a leer en el empaque o etiqueta como el nombre del artículo.
- Si es una tabla (Excel o tabla dentro de un Word/PDF), identifica cuál columna es \
el nombre/descripción del artículo y cuál es la cantidad — ignora columnas de precio, \
importe, IVA, etc.
- Ignora encabezados, totales, subtotales, precios, fechas, o cualquier texto que no \
sea un artículo en sí.
- Si de verdad no hay ningún artículo identificable, regresa una lista vacía.

Responde ÚNICAMENTE con JSON válido, sin texto antes ni después, con \
esta forma exacta:
{"items": [{"nombre": "texto tal cual lo leíste", "cantidad": 1}]}"""

PROMPT_LECTURA_PROMOCION = """Este archivo es un anuncio o volante de una PROMOCIÓN de productos \
(foto, captura de red social, imagen de diseño, PDF, etc.) — por ejemplo "Promoción de Septiembre", \
un combo, un paquete con descuento, 2x1, etc. Identifica:

1. Un nombre corto y fácil de identificar para esta promoción, basado en lo que dice el anuncio \
(ej. "Promoción Septiembre 2026", "Combo Higiene Dental", "2x1 Cepillos"). Si el anuncio no trae \
un nombre claro, invéntale uno corto y descriptivo basado en los artículos que incluye.
2. Cada artículo/producto incluido en la promoción, con su cantidad, y el PRECIO PROMOCIONAL \
(el precio especial/de oferta que se anuncia — NO el precio normal tachado, si se muestran ambos).

Reglas:
- Si no hay cantidad especificada para un artículo, usa 1.
- Si un artículo no tiene un precio promocional individual visible (ej. es un combo con un solo \
precio total para varios artículos), reparte el precio del combo entre los artículos lo mejor que \
puedas; si de plano no se puede repartir, usa 0 como precio_promocional para que se capture a mano.
- Ignora el precio normal/tachado si también aparece — solo interesa el precio de oferta.
- Si de verdad no hay ningún artículo identificable, regresa una lista vacía.

Responde ÚNICAMENTE con JSON válido, sin texto antes ni después, con esta forma exacta:
{"nombre_promocion": "texto sugerido", "items": [{"nombre": "texto tal cual lo leíste", "cantidad": 1, "precio_promocional": 0}]}"""

# Extensiones que tratamos como "documento de oficina" (se les extrae el
# texto/tabla en Python y se le manda a Claude como texto, no como imagen).
EXTENSIONES_DOCX = (".docx",)
EXTENSIONES_XLSX = (".xlsx", ".xlsm", ".xls")
EXTENSIONES_CSV = (".csv",)


def _api_key():
    key = os.getenv("ANTHROPIC_API_KEY", "").strip()
    if not key:
        raise RuntimeError(
            "Falta la variable de entorno ANTHROPIC_API_KEY. Configúrala en Render "
            "(Environment) con tu API key de console.anthropic.com para poder usar "
            "la lectura de imágenes/documentos."
        )
    return key


def _extraer_json(texto: str):
    """Claude normalmente responde solo el JSON, pero por si acaso viene
    con ```json ... ``` alrededor o algo de texto extra, lo limpiamos
    antes de parsear."""
    texto = texto.strip()
    texto = re.sub(r"^```json\s*|^```\s*|\s*```$", "", texto.strip(), flags=re.MULTILINE).strip()
    match = re.search(r"\{.*\}", texto, re.DOTALL)
    if match:
        texto = match.group(0)
    return json.loads(texto)


def _decodificar_base64(datos_base64: str, media_type_por_defecto: str):
    """Acepta tanto un data URL completo (data:...;base64,....) como puro
    base64. Regresa (bytes_crudos, media_type)."""
    media_type = media_type_por_defecto
    if datos_base64.startswith("data:"):
        try:
            media_type = datos_base64.split(";")[0].split(":")[1]
            datos_base64 = datos_base64.split(",", 1)[1]
        except (IndexError, ValueError):
            pass
    return base64.b64decode(datos_base64), media_type


def _texto_desde_docx(datos_bytes: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(datos_bytes))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            if any(celdas):
                partes.append(" | ".join(celdas))
    return "\n".join(partes)


def _texto_desde_xlsx(datos_bytes: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(datos_bytes), data_only=True)
    partes = []
    for hoja in wb.worksheets:
        for fila in hoja.iter_rows(values_only=True):
            valores = [str(v) for v in fila if v is not None and str(v).strip()]
            if valores:
                partes.append(" | ".join(valores))
    return "\n".join(partes)


def _texto_desde_csv(datos_bytes: bytes) -> str:
    for codificacion in ("utf-8", "latin-1"):
        try:
            return datos_bytes.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return datos_bytes.decode("utf-8", errors="ignore")


def _pedir_json_a_claude(bloques_contenido, empresa_id=None, system_extra="", tipo_consumo="leer_imagen_documento"):
    """Bajo nivel, compartido: arma la petición, la manda, maneja los
    errores comunes (401/429/etc.), registra el consumo, y regresa el
    dict ya parseado de JSON — sin saber ni validar la forma esperada,
    eso lo hace cada función de más alto nivel."""
    api_key = _api_key()
    body = {
        "model": MODELO_LECTURA_IMAGEN,
        "max_tokens": 8192,
        "system": (
            "Respondes ÚNICAMENTE con JSON válido — nada de texto antes, nada de texto después, "
            "nada de explicaciones, nada de marcado de código (```). Tu respuesta completa debe "
            "poder pasarse directo a json.loads() de Python sin ningún procesamiento previo."
            + system_extra
        ),
        "messages": [
            {"role": "user", "content": bloques_contenido},
        ],
    }
    try:
        r = requests.post(
            ANTHROPIC_API_URL,
            headers={
                "x-api-key": api_key,
                "anthropic-version": ANTHROPIC_VERSION,
                "content-type": "application/json",
            },
            json=body,
            timeout=90,
        )
    except requests.RequestException as e:
        raise RuntimeError(f"No se pudo conectar con la API de Claude: {e}")

    if r.status_code == 401:
        raise RuntimeError("La API key de Anthropic no es válida (revisa ANTHROPIC_API_KEY en Render).")
    if r.status_code == 429:
        raise RuntimeError("Se alcanzó el límite de uso de la API de Claude por ahora — intenta en un momento.")
    if not r.ok:
        detalle = r.text[:300]
        raise RuntimeError(f"La API de Claude respondió con error ({r.status_code}): {detalle}")

    data = r.json()
    try:
        uso = data.get("usage", {})
        tokens_in = uso.get("input_tokens", 0) or 0
        tokens_out = uso.get("output_tokens", 0) or 0
        costo = (tokens_in / 1_000_000 * 3.0) + (tokens_out / 1_000_000 * 15.0)  # precio aproximado, Sonnet
        db.registrar_consumo(empresa_id, tipo_consumo, tokens_in + tokens_out, "tokens", round(costo, 6))
    except Exception:
        pass
    if data.get("stop_reason") == "max_tokens":
        raise RuntimeError(
            "La lista de artículos es demasiado larga y la respuesta se cortó a medias. "
            "Intenta dividir el documento en partes más chicas."
        )
    bloques_texto = [b["text"] for b in data.get("content", []) if b.get("type") == "text"]
    texto_completo = "\n".join(bloques_texto)
    if not texto_completo.strip():
        raise RuntimeError("Claude no regresó contenido en la respuesta.")

    try:
        return _extraer_json(texto_completo)
    except (json.JSONDecodeError, ValueError):
        fragmento = texto_completo.strip().replace("\n", " ")[:200]
        raise RuntimeError(
            "No se pudo leer la respuesta de Claude como lista de artículos. "
            f"Esto fue lo que respondió: \"{fragmento}\""
        )


def _llamar_claude(bloques_contenido, empresa_id=None):
    parseado = _pedir_json_a_claude(
        bloques_contenido, empresa_id,
        system_extra=' Si no hay nada que reportar, responde exactamente: {"items": []}',
        tipo_consumo="leer_imagen_documento",
    )
    items = parseado.get("items", []) if isinstance(parseado, dict) else []
    if not isinstance(items, list):
        items = []
    resultado = []
    for it in items:
        if not isinstance(it, dict):
            continue
        nombre = (it.get("nombre") or "").strip()
        if not nombre:
            continue
        try:
            cantidad = float(it.get("cantidad") or 1)
        except (TypeError, ValueError):
            cantidad = 1
        resultado.append({"nombre": nombre, "cantidad": cantidad})
    return resultado


def _llamar_claude_promocion(bloques_contenido, empresa_id=None):
    parseado = _pedir_json_a_claude(
        bloques_contenido, empresa_id,
        system_extra=' Si no hay nada que reportar, responde exactamente: {"nombre_promocion": "", "items": []}',
        tipo_consumo="leer_promocion_imagen",
    )
    nombre_promocion = ""
    items_raw = []
    if isinstance(parseado, dict):
        nombre_promocion = (parseado.get("nombre_promocion") or "").strip()
        items_raw = parseado.get("items", [])
    if not isinstance(items_raw, list):
        items_raw = []
    resultado = []
    for it in items_raw:
        if not isinstance(it, dict):
            continue
        nombre = (it.get("nombre") or "").strip()
        if not nombre:
            continue
        try:
            cantidad = float(it.get("cantidad") or 1)
        except (TypeError, ValueError):
            cantidad = 1
        try:
            precio_promocional = float(it.get("precio_promocional") or 0)
        except (TypeError, ValueError):
            precio_promocional = 0
        resultado.append({"nombre": nombre, "cantidad": cantidad, "precio_promocional": precio_promocional})
    return {"nombre_promocion": nombre_promocion, "items": resultado}


def leer_lista_de_imagen(imagen_base64: str, media_type: str = "image/jpeg", empresa_id=None):
    """Compatibilidad con el nombre anterior — sigue funcionando igual
    que antes para imágenes."""
    return leer_lista_de_archivo(imagen_base64, nombre_archivo="imagen.jpg", media_type=media_type, empresa_id=empresa_id)


def leer_lista_de_archivo(datos_base64: str, nombre_archivo: str = "", media_type: str = "application/octet-stream", empresa_id=None):
    """Punto de entrada único: detecta si es imagen, PDF, Word, Excel o
    CSV, y arma el mensaje correcto para Claude en cada caso. Regresa
    una lista de dicts {"nombre": str, "cantidad": float}."""
    datos_bytes, media_type = _decodificar_base64(datos_base64, media_type)
    nombre_lower = (nombre_archivo or "").lower()

    if media_type.startswith("image/"):
        bloques = [
            {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": base64.b64encode(datos_bytes).decode()}},
            {"type": "text", "text": PROMPT_LECTURA_LISTA},
        ]
        return _llamar_claude(bloques, empresa_id)

    if media_type == "application/pdf" or nombre_lower.endswith(".pdf"):
        bloques = [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(datos_bytes).decode()}},
            {"type": "text", "text": PROMPT_LECTURA_LISTA},
        ]
        return _llamar_claude(bloques, empresa_id)

    if nombre_lower.endswith(EXTENSIONES_DOCX) or "wordprocessingml" in media_type:
        try:
            texto = _texto_desde_docx(datos_bytes)
        except Exception as e:
            raise RuntimeError(f"No se pudo leer el archivo Word: {e}")
        return _llamar_claude_con_texto(texto, empresa_id)

    if nombre_lower.endswith(EXTENSIONES_XLSX) or "spreadsheetml" in media_type or media_type == "application/vnd.ms-excel":
        try:
            texto = _texto_desde_xlsx(datos_bytes)
        except Exception as e:
            raise RuntimeError(f"No se pudo leer el archivo Excel: {e}")
        return _llamar_claude_con_texto(texto, empresa_id)

    if nombre_lower.endswith(EXTENSIONES_CSV) or media_type == "text/csv":
        texto = _texto_desde_csv(datos_bytes)
        return _llamar_claude_con_texto(texto, empresa_id)

    raise RuntimeError(
        "Ese tipo de archivo no lo puedo leer todavía. Sube una imagen (foto), PDF, Word (.docx) o Excel (.xlsx/.csv)."
    )


def _llamar_claude_con_texto(texto_extraido: str, empresa_id=None):
    if not texto_extraido.strip():
        return []
    # Si el documento es enorme, lo recortamos para no disparar el costo/tiempo
    # de una sola llamada — para listas de artículos esto es más que suficiente.
    texto_extraido = texto_extraido[:40000]
    bloques = [
        {"type": "text", "text": f"Contenido del documento:\n\n{texto_extraido}\n\n{PROMPT_LECTURA_LISTA}"},
    ]
    return _llamar_claude(bloques, empresa_id)


def leer_promocion_de_archivo(datos_base64: str, nombre_archivo: str = "", media_type: str = "application/octet-stream", empresa_id=None):
    """Lee un anuncio/volante de una promoción (imagen o PDF) y regresa
    {"nombre_promocion": str, "items": [{"nombre", "cantidad", "precio_promocional"}]}.
    A diferencia de leer_lista_de_archivo (para cotizaciones), aquí también
    se extrae un nombre sugerido para la promoción y el precio de OFERTA de
    cada artículo — no el de Microsip. Solo imagen/PDF: un anuncio de
    promoción es por naturaleza visual (no tiene sentido para Excel/CSV)."""
    datos_bytes, media_type = _decodificar_base64(datos_base64, media_type)
    nombre_lower = (nombre_archivo or "").lower()

    if media_type.startswith("image/"):
        bloques = [
            {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": base64.b64encode(datos_bytes).decode()}},
            {"type": "text", "text": PROMPT_LECTURA_PROMOCION},
        ]
        return _llamar_claude_promocion(bloques, empresa_id)

    if media_type == "application/pdf" or nombre_lower.endswith(".pdf"):
        bloques = [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(datos_bytes).decode()}},
            {"type": "text", "text": PROMPT_LECTURA_PROMOCION},
        ]
        return _llamar_claude_promocion(bloques, empresa_id)

    raise RuntimeError(
        "Para leer una promoción sube una imagen (foto o captura) o un PDF del anuncio."
    )
