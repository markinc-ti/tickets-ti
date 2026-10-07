"""Asistencia (checador ZKTeco / BioTime Pro) para RH.

- Las checadas llegan de dos formas:
  1) El programa biotime_sync.ps1 que corre en la PC de BioTime (cada 5 min
     lee la API local de BioTime y las manda a /api/rh/checadas/sync).
  2) Importando el archivo que se exporta de BioTime (Excel o CSV).
- Con el horario de cada empleado se calculan entrada, salida, horas,
  retardos y faltas (las faltas con incidencia aprobada o vacaciones salen
  como justificadas).
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime, time, timedelta

HORARIO_DEFAULT = {"hora_entrada": "09:00", "hora_salida": "18:00", "dias": "1,2,3,4,5,6", "tolerancia_min": 10}
NOMBRES_DIA = ["", "Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]


def normalizar_codigo(codigo) -> str:
    """Número de empleado comparable: sin espacios ni ceros a la izquierda."""
    c = str(codigo or "").strip()
    if re.fullmatch(r"\d+(\.0+)?", c):
        c = c.split(".")[0].lstrip("0") or "0"
    return c.upper()


def _sin_acentos(t: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFD", str(t or "")) if unicodedata.category(ch) != "Mn").lower().strip()


# ---------- Fechas ----------

def a_datetime(valor, fecha_aparte=None):
    """Convierte lo que venga del archivo/API a datetime. Formatos de México
    (día/mes/año) cuando hay ambigüedad."""
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None, microsecond=0)
    if isinstance(valor, time) and fecha_aparte is not None:
        f = a_fecha(fecha_aparte)
        return datetime.combine(f, valor.replace(microsecond=0)) if f else None
    texto = str(valor or "").strip()
    if fecha_aparte is not None and texto:
        f = a_fecha(fecha_aparte)
        h = _a_hora(texto)
        return datetime.combine(f, h) if f and h else None
    if not texto:
        return None
    texto = texto.replace("T", " ").split("+")[0].split(".")[0].strip()
    texto = re.sub(r"\s*(a\.?\s?m\.?|p\.?\s?m\.?)$", lambda m: " " + ("PM" if "p" in m.group(1).lower() else "AM"), texto, flags=re.I)
    formatos = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M",
                "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
                "%d/%m/%Y %I:%M:%S %p", "%d/%m/%Y %I:%M %p", "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %I:%M %p",
                "%d/%m/%y %H:%M", "%d/%m/%y %H:%M:%S"]
    for fmt in formatos:
        try:
            return datetime.strptime(texto, fmt)
        except ValueError:
            continue
    return None


def a_fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip().split(" ")[0].split("T")[0]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    return None


def _a_hora(valor):
    if isinstance(valor, time):
        return valor.replace(microsecond=0)
    texto = str(valor or "").strip().upper()
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p"):
        try:
            return datetime.strptime(texto, fmt).time()
        except ValueError:
            continue
    return None


def a_minutos(hhmm: str) -> int:
    h, m = (hhmm or "00:00").split(":")[:2]
    return int(h) * 60 + int(m)


# ---------- Importar archivo de BioTime ----------

def _rol_columna(encabezado: str):
    e = _sin_acentos(encabezado)
    e = re.sub(r"\s+", " ", e)
    if not e:
        return None
    if e in ("punch time", "hora de marcaje", "fecha y hora", "fecha/hora", "fecha hora", "marcaje",
             "datetime", "date time", "hora de checada", "checada", "fecha de marcaje", "hora de registro") \
            or ("punch" in e and "time" in e) or ("marcaje" in e and "hora" in e):
        return "fecha_hora"
    if e in ("date", "fecha", "dia", "fecha de checada"):
        return "fecha"
    if e in ("time", "hora"):
        return "hora"
    if e in ("employee id", "emp code", "emp_code", "personnel id", "id de empleado", "id empleado",
             "no. empleado", "no empleado", "numero de empleado", "num. empleado", "codigo", "codigo de empleado",
             "id de personal", "id personal", "id", "no.", "ac-no.", "ac no", "user id", "id de usuario") \
            or ("empleado" in e and ("id" in e or "num" in e or "no" in e.split())) or ("personnel" in e and "id" in e):
        return "codigo"
    if e in ("first name", "nombre", "nombre(s)", "name", "nombre completo", "full name"):
        return "nombre"
    if e in ("last name", "apellido", "apellidos"):
        return "apellido"
    if e in ("punch state", "estado de marcaje", "estado", "state", "tipo de marcaje", "tipo", "estado de checada"):
        return "tipo"
    if e in ("device name", "terminal", "dispositivo", "nombre del dispositivo", "terminal alias", "device",
             "alias del dispositivo", "serial number", "numero de serie", "terminal sn"):
        return "dispositivo"
    if e in ("area", "area name", "nombre del area", "area alias", "sucursal"):
        return "area"
    return None


def _filas_de_archivo(datos: bytes, nombre: str):
    nombre = (nombre or "").lower()
    if datos[:2] == b"PK" or nombre.endswith((".xlsx", ".xlsm")):
        import openpyxl
        libro = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
        for hoja in libro.worksheets:
            filas = [list(r) for r in hoja.iter_rows(values_only=True)]
            if any(any(c not in (None, "") for c in f) for f in filas):
                yield filas
        return
    if datos[:4] == b"\xd0\xcf\x11\xe0" or nombre.endswith(".xls"):
        raise ValueError("Ese archivo es de Excel viejo (.xls). En BioTime expórtalo como Excel (.xlsx) o CSV.")
    texto = None
    for cod in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            texto = datos.decode(cod)
            if "\x00" not in texto:
                break
        except UnicodeDecodeError:
            continue
    muestra = texto[:4000]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    yield [list(r) for r in csv.reader(io.StringIO(texto), dialecto)]


def leer_archivo_biotime(datos: bytes, nombre: str) -> list:
    """Regresa [{codigo, nombre, fecha_hora (datetime), tipo, dispositivo, area}]."""
    checadas = []
    hubo_tabla = False
    for filas in _filas_de_archivo(datos, nombre):
        # El encabezado puede no estar en la primera fila (los reportes traen título)
        idx, roles = None, None
        for i, fila in enumerate(filas[:25]):
            r = [_rol_columna(str(c or "")) for c in fila]
            if "codigo" in r and ("fecha_hora" in r or ("fecha" in r and "hora" in r)):
                idx, roles = i, r
                break
        if idx is None:
            continue
        hubo_tabla = True
        col = {}
        for j, rol in enumerate(roles):
            if rol and rol not in col:
                col[rol] = j
        for fila in filas[idx + 1:]:
            get = lambda k: fila[col[k]] if k in col and col[k] < len(fila) else None
            codigo = normalizar_codigo(get("codigo"))
            if not codigo or codigo in ("0", "NONE"):
                continue
            if "fecha_hora" in col:
                fh = a_datetime(get("fecha_hora"))
            else:
                fh = a_datetime(get("hora"), fecha_aparte=get("fecha"))
            if not fh:
                continue
            nombre_c = " ".join(str(x).strip() for x in (get("nombre"), get("apellido")) if x not in (None, ""))
            checadas.append({
                "codigo": codigo, "nombre": nombre_c or None, "fecha_hora": fh,
                "tipo": (str(get("tipo")).strip() or None) if get("tipo") not in (None, "") else None,
                "dispositivo": (str(get("dispositivo")).strip() or None) if get("dispositivo") not in (None, "") else None,
                "area": (str(get("area")).strip() or None) if get("area") not in (None, "") else None,
            })
    if not hubo_tabla:
        raise ValueError("No encontré las columnas del reporte (se necesitan al menos el número de empleado y la fecha/hora de la checada). "
                         "En BioTime usa Asistencia → Transacciones → Exportar.")
    return checadas


def checada_de_api(t: dict):
    """Una transacción de la API de BioTime (o la que manda el programa de la PC)."""
    codigo = normalizar_codigo(t.get("codigo") or t.get("emp_code"))
    fh = a_datetime(t.get("fecha_hora") or t.get("punch_time"))
    if not codigo or not fh:
        return None
    nombre = t.get("nombre") or " ".join(x for x in (t.get("first_name"), t.get("last_name")) if x) or None
    return {
        "codigo": codigo, "nombre": (str(nombre).strip() or None) if nombre else None, "fecha_hora": fh,
        "tipo": t.get("tipo") or t.get("punch_state_display") or t.get("punch_state"),
        "dispositivo": t.get("dispositivo") or t.get("terminal_alias") or t.get("terminal_sn"),
        "area": t.get("area") or t.get("area_alias"),
        "biotime_id": str(t.get("id")) if t.get("id") is not None else None,
    }


# ---------- Reporte ----------

def _rangos_justificados(lista):
    """[(desde, hasta, motivo)] → función fecha → motivo|None"""
    def buscar(d):
        for a, b, motivo in lista:
            if a and a <= d <= (b or a):
                return motivo
        return None
    return buscar


def calcular_reporte(empleados, checadas, horarios, horario_default, desde: date, hasta: date, hoy: date, justificaciones,
                     inicio_datos: date | None = None):
    """
    empleados: [{id, nombre, numero_empleado, sucursal}]
    checadas: [{codigo, fecha_hora(datetime), dispositivo}]
    horarios: {usuario_id: {hora_entrada, hora_salida, dias, tolerancia_min}}
    justificaciones: {usuario_id: [(desde, hasta, motivo)]}
    inicio_datos: primer día con checadas en la empresa; antes de eso no se cuentan faltas
    """
    por_codigo = {}
    for c in checadas:
        por_codigo.setdefault(c["codigo"], []).append(c)
    codigos_empleados = {normalizar_codigo(e["numero_empleado"]): e for e in empleados if e.get("numero_empleado")}
    sin_empleado = {}
    for codigo, lista in por_codigo.items():
        if codigo not in codigos_empleados:
            sin_empleado[codigo] = {"codigo": codigo, "nombre": next((x.get("nombre") for x in lista if x.get("nombre")), None),
                                    "checadas": len(lista)}

    resultado = []
    for e in empleados:
        cod = normalizar_codigo(e.get("numero_empleado"))
        h = {**horario_default, **(horarios.get(e["id"]) or {})}
        dias_lab = {int(x) for x in str(h["dias"]).split(",") if x.strip().isdigit()}
        tol = int(h.get("tolerancia_min") or 0)
        ent_min, sal_min = a_minutos(h["hora_entrada"]), a_minutos(h["hora_salida"])
        justif = _rangos_justificados(justificaciones.get(e["id"], []))
        del_empleado = {}
        for c in por_codigo.get(cod, []) if cod else []:
            del_empleado.setdefault(c["fecha_hora"].date(), []).append(c)
        dias = []
        resumen = {"dias_trabajados": 0, "retardos": 0, "minutos_retardo": 0, "faltas": 0, "justificadas": 0,
                   "sin_salida": 0, "minutos_trabajados": 0}
        d = desde
        while d <= hasta:
            lista = sorted(del_empleado.get(d, []), key=lambda x: x["fecha_hora"])
            laborable = d.isoweekday() in dias_lab
            dia = {"fecha": d.isoformat(), "dia": NOMBRES_DIA[d.isoweekday()], "laborable": laborable,
                   "checadas": [x["fecha_hora"].strftime("%H:%M") for x in lista],
                   "entrada": None, "salida": None, "minutos": 0, "estatus": None, "minutos_retardo": 0, "motivo": None}
            if lista:
                entrada = lista[0]["fecha_hora"]
                dia["entrada"] = entrada.strftime("%H:%M")
                resumen["dias_trabajados"] += 1
                if len(lista) > 1:
                    salida = lista[-1]["fecha_hora"]
                    dia["salida"] = salida.strftime("%H:%M")
                    dia["minutos"] = int((salida - entrada).total_seconds() // 60)
                    resumen["minutos_trabajados"] += dia["minutos"]
                minutos_entrada = entrada.hour * 60 + entrada.minute
                if laborable and minutos_entrada > ent_min + tol:
                    dia["estatus"] = "retardo"
                    dia["minutos_retardo"] = minutos_entrada - ent_min
                    resumen["retardos"] += 1
                    resumen["minutos_retardo"] += dia["minutos_retardo"]
                elif not laborable:
                    dia["estatus"] = "dia_no_laboral"
                else:
                    dia["estatus"] = "puntual"
                if len(lista) == 1 and d < hoy:
                    dia["estatus"] = dia["estatus"] if dia["estatus"] == "retardo" else "sin_salida"
                    resumen["sin_salida"] += 1
            elif inicio_datos and d < inicio_datos:
                dia["estatus"] = None
            elif laborable and d < hoy:
                motivo = justif(d)
                if motivo:
                    dia["estatus"], dia["motivo"] = "justificada", motivo
                    resumen["justificadas"] += 1
                else:
                    dia["estatus"] = "falta"
                    resumen["faltas"] += 1
            elif laborable and d == hoy:
                dia["estatus"] = "pendiente" if datetime.now().hour * 60 < sal_min else "falta"
                if dia["estatus"] == "falta":
                    motivo = justif(d)
                    if motivo:
                        dia["estatus"], dia["motivo"] = "justificada", motivo
                        resumen["justificadas"] += 1
                    else:
                        resumen["faltas"] += 1
            else:
                dia["estatus"] = "descanso" if not laborable else None
            dias.append(dia)
            d += timedelta(days=1)
        resultado.append({
            "usuario_id": e["id"], "nombre": e["nombre"], "numero_empleado": e.get("numero_empleado"),
            "sucursal": e.get("sucursal"), "horario": {**h, "tolerancia_min": tol}, "tiene_horario_propio": e["id"] in horarios,
            "resumen": resumen, "dias": dias,
        })
    resultado.sort(key=lambda r: (-(r["resumen"]["faltas"] + r["resumen"]["retardos"]), r["nombre"] or ""))
    return {"empleados": resultado, "codigos_sin_empleado": sorted(sin_empleado.values(), key=lambda x: x["codigo"])}


NOMBRES_ESTATUS = {"puntual": "Puntual", "retardo": "Retardo", "falta": "Falta", "justificada": "Justificada",
                   "sin_salida": "Sin salida", "dia_no_laboral": "Día no laboral", "descanso": "Descanso", "pendiente": "Pendiente"}


def _hhmm(minutos: int) -> str:
    return f"{minutos // 60}:{minutos % 60:02d}"


def excel_reporte(reporte: dict, desde: date, hasta: date, empresa: str) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    libro = openpyxl.Workbook()
    negrita = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="D8192F")
    colores = {"retardo": "FFF2CC", "falta": "F8CBAD", "justificada": "DDEBF7", "sin_salida": "FCE4D6"}

    hoja = libro.active
    hoja.title = "Resumen"
    hoja.append([f"Asistencia {empresa} — del {desde.isoformat()} al {hasta.isoformat()}"])
    hoja["A1"].font = Font(bold=True, size=13)
    hoja.append([])
    enc = ["Empleado", "No. empleado", "Sucursal", "Horario", "Días trabajados", "Retardos", "Minutos de retardo",
           "Faltas", "Justificadas", "Sin salida", "Horas trabajadas"]
    hoja.append(enc)
    for c in hoja[3]:
        c.font, c.fill = negrita, relleno
    for e in reporte["empleados"]:
        r, h = e["resumen"], e["horario"]
        hoja.append([e["nombre"], e["numero_empleado"], e.get("sucursal") or "",
                     f"{h['hora_entrada']}-{h['hora_salida']} (tol. {h['tolerancia_min']} min)",
                     r["dias_trabajados"], r["retardos"], r["minutos_retardo"], r["faltas"], r["justificadas"],
                     r["sin_salida"], _hhmm(r["minutos_trabajados"])])
    for col, ancho in zip("ABCDEFGHIJK", (32, 13, 18, 26, 15, 10, 17, 8, 12, 11, 15)):
        hoja.column_dimensions[col].width = ancho

    det = libro.create_sheet("Detalle")
    enc = ["Empleado", "No. empleado", "Fecha", "Día", "Entrada", "Salida", "Horas", "Estatus", "Min. retardo", "Motivo", "Todas las checadas"]
    det.append(enc)
    for c in det[1]:
        c.font, c.fill = negrita, relleno
    for e in reporte["empleados"]:
        for d in e["dias"]:
            if not d["estatus"] or d["estatus"] == "descanso":
                continue
            det.append([e["nombre"], e["numero_empleado"], d["fecha"], d["dia"], d["entrada"] or "", d["salida"] or "",
                        _hhmm(d["minutos"]) if d["minutos"] else "", NOMBRES_ESTATUS.get(d["estatus"], d["estatus"]),
                        d["minutos_retardo"] or "", d["motivo"] or "", ", ".join(d["checadas"])])
            color = colores.get(d["estatus"])
            if color:
                for c in det[det.max_row]:
                    c.fill = PatternFill("solid", fgColor=color)
    for col, ancho in zip("ABCDEFGHIJK", (32, 13, 12, 6, 9, 9, 8, 14, 12, 30, 40)):
        det.column_dimensions[col].width = ancho
    for hoja_ in (hoja, det):
        hoja_.freeze_panes = "A4" if hoja_ is hoja else "A2"
        for fila in hoja_.iter_rows():
            for c in fila:
                c.alignment = Alignment(vertical="center")
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ---------- Programa para la PC de BioTime ----------

PLANTILLA_AGENTE = r'''# biotime_sync.ps1 -- Manda las checadas de BioTime Pro a tickets-ti.
# Se instala en la PC donde esta BioTime. Corre cada 5 minutos con el
# Programador de tareas de Windows (ver "Instalar" abajo). No hay que
# abrir puertos del router: esta PC es la que se conecta a internet.
#
# Instalar (PowerShell como Administrador, en la carpeta de este archivo):
#   1) Escribe abajo el usuario y contrasena de BioTime (los de entrar al
#      BioTime por el navegador) y, si no es el puerto 80, la direccion.
#   2) Prueba:      powershell -ExecutionPolicy Bypass -File .\biotime_sync.ps1
#   3) Programalo:  schtasks /Create /F /SC MINUTE /MO 5 /RU SYSTEM /TN "BioTime a tickets-ti" /TR "powershell -NoProfile -ExecutionPolicy Bypass -File \"%RUTA%\""
#      (cambia %RUTA% por la ruta completa de este archivo, ej. C:\biotime_sync\biotime_sync.ps1)

# ---- Configuracion de esta PC ----
$BioTimeUrl      = "http://127.0.0.1:80"      # como entras a BioTime en el navegador de esta PC (ej. http://127.0.0.1:8081)
$BioTimeUsuario  = "admin"
$BioTimePassword = "CAMBIAR"
$Equipo          = $env:COMPUTERNAME           # nombre con el que se ve esta PC en tickets-ti

# ---- No cambiar (ya viene de tickets-ti) ----
$TicketsUrl   = "__TICKETS_URL__"
$TicketsToken = "__TOKEN__"
$DiasAtras    = 3      # la primera vez manda los ultimos 3 dias; despues lo nuevo (repasa 1 dia por si un checador subio tarde)
# -------------------------------------------

$ErrorActionPreference = "Stop"
[Net.ServicePointManager]::SecurityProtocol = [Net.SecurityProtocolType]::Tls12
$Carpeta = Split-Path -Parent $MyInvocation.MyCommand.Path
if (-not $Carpeta) { $Carpeta = (Get-Location).Path }
$Estado  = Join-Path $Carpeta "biotime_sync_estado.txt"
$Bitacora = Join-Path $Carpeta "biotime_sync_bitacora.txt"

function Escribir($texto) {
    $linea = (Get-Date -Format "yyyy-MM-dd HH:mm:ss") + "  " + $texto
    Write-Host $linea
    Add-Content -Path $Bitacora -Value $linea -Encoding UTF8
}

function Get-TokenBioTime {
    $cuerpo = @{ username = $BioTimeUsuario; password = $BioTimePassword } | ConvertTo-Json
    try {
        $r = Invoke-RestMethod -Method Post -Uri ($BioTimeUrl.TrimEnd("/") + "/jwt-api-token-auth/") -ContentType "application/json" -Body $cuerpo
        return "JWT " + $r.token
    } catch {
        $r = Invoke-RestMethod -Method Post -Uri ($BioTimeUrl.TrimEnd("/") + "/api-token-auth/") -ContentType "application/json" -Body $cuerpo
        return "Token " + $r.token
    }
}

try {
    if ($BioTimePassword -eq "CAMBIAR") { throw "Falta poner el usuario y contrasena de BioTime al inicio de este archivo." }
    $ahora = Get-Date
    $desde = $ahora.AddDays(-$DiasAtras)
    if (Test-Path $Estado) {
        $guardado = (Get-Content $Estado -Raw).Trim()
        if ($guardado) { $desde = [datetime]::ParseExact($guardado, "yyyy-MM-dd HH:mm:ss", $null).AddDays(-1) }
    }
    $auth = Get-TokenBioTime
    $inicio = [uri]::EscapeDataString($desde.ToString("yyyy-MM-dd HH:mm:ss"))
    $fin = [uri]::EscapeDataString($ahora.ToString("yyyy-MM-dd HH:mm:ss"))
    $url = $BioTimeUrl.TrimEnd("/") + "/iclock/api/transactions/?page_size=1000&start_time=$inicio&end_time=$fin"
    $checadas = New-Object System.Collections.ArrayList
    $paginas = 0
    while ($url -and $paginas -lt 200) {
        $r = Invoke-RestMethod -Method Get -Uri $url -Headers @{ Authorization = $auth }
        $lista = $r.data
        if ($null -eq $lista) { $lista = $r.results }
        foreach ($t in $lista) {
            [void]$checadas.Add(@{
                id = $t.id; codigo = [string]$t.emp_code
                nombre = (([string]$t.first_name + " " + [string]$t.last_name).Trim())
                fecha_hora = [string]$t.punch_time; tipo = [string]$t.punch_state_display
                dispositivo = [string]$(if ($t.terminal_alias) { $t.terminal_alias } else { $t.terminal_sn })
                area = [string]$t.area_alias
            })
        }
        $url = $r.next
        $paginas++
    }
    $enviadas = 0; $nuevas = 0
    for ($i = 0; $i -lt $checadas.Count; $i += 2000) {
        $fin_lote = [Math]::Min($i + 2000, $checadas.Count) - 1
        $lote = @($checadas[$i..$fin_lote])
        $json = @{ equipo = $Equipo; checadas = $lote } | ConvertTo-Json -Depth 4 -Compress
        $bytes = [System.Text.Encoding]::UTF8.GetBytes($json)
        $r = Invoke-RestMethod -Method Post -Uri ($TicketsUrl.TrimEnd("/") + "/api/rh/checadas/sync") -ContentType "application/json; charset=utf-8" -Headers @{ "X-Token-Checador" = $TicketsToken } -Body $bytes
        $enviadas += $lote.Count; $nuevas += [int]$r.nuevas
    }
    if ($checadas.Count -eq 0) {
        $json = @{ equipo = $Equipo; checadas = @() } | ConvertTo-Json -Compress
        $null = Invoke-RestMethod -Method Post -Uri ($TicketsUrl.TrimEnd("/") + "/api/rh/checadas/sync") -ContentType "application/json; charset=utf-8" -Headers @{ "X-Token-Checador" = $TicketsToken } -Body ([System.Text.Encoding]::UTF8.GetBytes($json))
    }
    Set-Content -Path $Estado -Value $ahora.ToString("yyyy-MM-dd HH:mm:ss") -Encoding ASCII
    Escribir ("OK: " + $enviadas + " checadas revisadas, " + $nuevas + " nuevas en tickets-ti.")
} catch {
    Escribir ("ERROR: " + $_.Exception.Message)
    exit 1
}
'''


def script_agente(url_tickets: str, token: str) -> str:
    return PLANTILLA_AGENTE.replace("__TICKETS_URL__", url_tickets).replace("__TOKEN__", token).replace("\n", "\r\n")
